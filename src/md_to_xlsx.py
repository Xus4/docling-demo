"""从 Markdown 中提取 GFM/HTML 表格并写入 Excel（每个表格一个 Sheet）。"""

from __future__ import annotations

import html as html_mod
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.table_semantic.table_blocks import iter_table_blocks

_HTML_TAG_RE = re.compile(r"<[^>]+>", re.DOTALL)
_TR_RE = re.compile(r"<tr[^>]*>(.*?)</tr>", re.IGNORECASE | re.DOTALL)
_TD_RE = re.compile(r"<(td|th)[^>]*>(.*?)</\1>", re.IGNORECASE | re.DOTALL)
_SEP_RE = re.compile(r"^:?-{3,}:?$")

_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_VERTICAL_ALIGN = Alignment(wrap_text=True, vertical="center")


def _parse_gfm_rows(raw: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in raw.strip().splitlines():
        s = line.strip()
        if not s or "|" not in s:
            continue
        cells = [c.strip() for c in s.split("|")]
        if cells and cells[0] == "":
            cells.pop(0)
        if cells and cells[-1] == "":
            cells.pop()
        if all(_SEP_RE.match(c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _parse_html_rows(raw: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for tr_m in _TR_RE.finditer(raw):
        cells: list[str] = []
        for td_m in _TD_RE.finditer(tr_m.group(1)):
            content = _HTML_TAG_RE.sub("", td_m.group(2))
            content = html_mod.unescape(content).strip()
            cells.append(content)
        if cells:
            rows.append(cells)
    return rows


def _pad_rows(rows: list[list[str]]) -> list[list[str]]:
    if not rows:
        return rows
    max_c = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_c:
            r.append("")
    return rows


def md_to_xlsx(md_path: Path, xlsx_path: Path) -> bool:
    text = md_path.read_text(encoding="utf-8")
    blocks = list(iter_table_blocks(text))
    if not blocks:
        return False

    wb = Workbook()
    wb.remove(wb.active)

    for idx, block in enumerate(blocks):
        rows = _parse_gfm_rows(block.raw) if block.kind == "gfm_table" else _parse_html_rows(block.raw)
        rows = _pad_rows(rows)
        if not rows:
            continue

        ws = wb.create_sheet(title=f"表格{idx + 1}")
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = ws.cell(row=r_idx + 1, column=c_idx + 1, value=val)
                cell.alignment = _VERTICAL_ALIGN
                cell.border = _THIN_BORDER
                if r_idx == 0:
                    cell.font = _HEADER_FONT
                    cell.fill = _HEADER_FILL

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(xlsx_path))
    return True
