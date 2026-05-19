"""MD 表格 → Excel（openpyxl）。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.md_to_xlsx import md_to_xlsx


class TestMdToXlsx(unittest.TestCase):
    def test_gfm_table_basic(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text("| 姓名 | 年龄 |\n| --- | --- |\n| 张三 | 28 |\n| 李四 | 32 |\n", encoding="utf-8")
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            self.assertTrue(xlsx_path.is_file())
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A1"].value, "姓名")
            self.assertEqual(ws["B1"].value, "年龄")
            self.assertEqual(ws["A2"].value, "张三")
            self.assertEqual(ws["B2"].value, "28")
            self.assertEqual(ws["A3"].value, "李四")
            self.assertEqual(ws["B3"].value, "32")

    def test_gfm_table_alignment_colons(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text("| 产品 | 价格 |\n| :--- | ---: |\n| A | 100 |\n", encoding="utf-8")
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A1"].value, "产品")
            self.assertEqual(ws["B1"].value, "价格")
            self.assertEqual(ws["A2"].value, "A")
            self.assertEqual(ws["B2"].value, "100")

    def test_html_table_basic(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text(
                "<table><thead><tr><th>城市</th><th>人口</th></tr></thead>"
                "<tbody><tr><td>北京</td><td>2154万</td></tr></tbody></table>\n",
                encoding="utf-8",
            )
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A1"].value, "城市")
            self.assertEqual(ws["B1"].value, "人口")
            self.assertEqual(ws["A2"].value, "北京")
            self.assertEqual(ws["B2"].value, "2154万")

    def test_no_table_returns_false(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text("# 标题\n\n只有文字，没有表格。\n", encoding="utf-8")
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertFalse(ok)

    def test_multiple_tables_multiple_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text(
                "| A | B |\n| --- | --- |\n| 1 | 2 |\n\n"
                "<table><tr><td>X</td><td>Y</td></tr><tr><td>3</td><td>4</td></tr></table>\n",
                encoding="utf-8",
            )
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            self.assertEqual(wb.sheetnames, ["表格1", "表格2"])
            ws1 = wb["表格1"]
            self.assertEqual(ws1["A1"].value, "A")
            self.assertEqual(ws1["B1"].value, "B")
            ws2 = wb["表格2"]
            self.assertEqual(ws2["A1"].value, "X")
            self.assertEqual(ws2["B1"].value, "Y")

    def test_skips_table_inside_fence(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text(
                "```\n| a | b |\n| --- | --- |\n| 1 | 2 |\n```\n\n"
                "| x | y |\n| --- | --- |\n| 3 | 4 |\n",
                encoding="utf-8",
            )
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A1"].value, "x")
            self.assertEqual(ws["B1"].value, "y")

    def test_empty_row_padding(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text("| A | B | C |\n| --- | --- | --- |\n| 1 |\n", encoding="utf-8")
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A2"].value, "1")
            self.assertIsNone(ws["B2"].value)

    def test_html_entities_decoded(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text(
                "<table><tr><td>&amp;</td><td>&lt;3</td><td>&#x2603;</td></tr></table>",
                encoding="utf-8",
            )
            xlsx_path = Path(td) / "test.xlsx"
            ok = md_to_xlsx(md_path, xlsx_path)
            self.assertTrue(ok)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertEqual(ws["A1"].value, "&")
            self.assertEqual(ws["B1"].value, "<3")
            self.assertEqual(ws["C1"].value, "\u2603")

    def test_header_style_applied(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            md_path = Path(td) / "test.md"
            md_path.write_text("| Key | Value |\n| --- | --- |\n| 1 | 2 |\n", encoding="utf-8")
            xlsx_path = Path(td) / "test.xlsx"
            md_to_xlsx(md_path, xlsx_path)
            wb = load_workbook(xlsx_path)
            ws = wb["表格1"]
            self.assertTrue(ws["A1"].font.bold)
            self.assertFalse(ws["A2"].font.bold)


if __name__ == "__main__":
    unittest.main()
